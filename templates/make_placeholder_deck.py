#!/usr/bin/env python3
"""Generate a NEUTRAL placeholder deck (templates/Template.pptx).

This is a generic, brand-free skeleton so the "deck mode" has something concrete to clone. Swap it for your
OWN branded .pptx (same filename) and the deck builder will use yours instead. The generator is included so
the binary isn't opaque in a public repo — you can see exactly what's in it and regenerate it.

    python3 templates/make_placeholder_deck.py
"""
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

OUT = Path(__file__).resolve().parent / "Template.pptx"

# a neutral placeholder accent (slate blue-grey) — NOT any company's brand color
ACCENT = RGBColor(0x3B, 0x4A, 0x5A)
MUTED = RGBColor(0x6B, 0x72, 0x80)


def _set(tf, text, size, color=None, bold=False):
    tf.text = text
    p = tf.paragraphs[0]
    p.font.size = Pt(size)
    p.font.bold = bold
    if color is not None:
        p.font.color.rgb = color


def build() -> Path:
    prs = Presentation()                      # default 10 x 7.5 in
    prs.slide_width = Inches(13.333)          # 16:9
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]              # fully blank layout

    # --- Slide 1: title ---
    s = prs.slides.add_slide(blank)
    tb = s.shapes.add_textbox(Inches(0.9), Inches(2.6), Inches(11.5), Inches(1.4))
    _set(tb.text_frame, "Your Deck Template", 44, ACCENT, bold=True)
    sub = s.shapes.add_textbox(Inches(0.9), Inches(4.0), Inches(11.5), Inches(1.0))
    _set(sub.text_frame, "Replace this file with your own branded .pptx (keep the filename Template.pptx).",
         18, MUTED)

    # --- Slide 2: takeaway-title content example ---
    s = prs.slides.add_slide(blank)
    t = s.shapes.add_textbox(Inches(0.9), Inches(0.6), Inches(11.5), Inches(1.0))
    _set(t.text_frame, "One takeaway per slide — put the point in the title", 28, ACCENT, bold=True)
    body = s.shapes.add_textbox(Inches(0.9), Inches(2.0), Inches(11.5), Inches(4.5))
    tf = body.text_frame
    tf.word_wrap = True
    _set(tf, "• The chart/number goes here (the deck builder inserts it).", 18)
    for line in ("• Say what it means, not just what it shows.",
                 "• State the caveat plainly.",
                 "• This is a generic placeholder — no real data, no branding."):
        p = tf.add_paragraph()
        p.text = line
        p.font.size = Pt(18)

    # scrub metadata so the placeholder carries no author/company info
    cp = prs.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = "Deck Template (placeholder)"
    cp.comments = "Generic brand-free placeholder — replace with your own .pptx."

    prs.save(str(OUT))
    return OUT


if __name__ == "__main__":
    p = build()
    print(f"wrote {p}  ({p.stat().st_size:,} bytes)")
